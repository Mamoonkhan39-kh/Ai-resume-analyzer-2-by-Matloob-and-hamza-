import json
import os
import re
from typing import Any

import streamlit as st
from docx import Document
from google import genai
from pypdf import PdfReader

st.set_page_config(page_title="Resume ATS & Interview Coach", page_icon="📄", layout="wide")

MAX_FILE_MB = 10
MAX_TEXT_CHARS = 45_000
DEFAULT_MODEL = "gemini-2.5-flash"


def get_api_key() -> str | None:
    """Read the Gemini key from Streamlit secrets first, then environment."""
    try:
        key = st.secrets.get("GEMINI_API_KEY", "")
    except Exception:
        key = ""
    return key or os.getenv("GEMINI_API_KEY")


def extract_resume_text(uploaded_file) -> str:
    """Extract text from PDF, DOCX, or TXT uploads."""
    name = uploaded_file.name.lower()
    raw = uploaded_file.getvalue()
    if len(raw) > MAX_FILE_MB * 1024 * 1024:
        raise ValueError(f"File is too large. Please upload a file under {MAX_FILE_MB} MB.")

    if name.endswith(".pdf"):
        from io import BytesIO
        reader = PdfReader(BytesIO(raw))
        pages = []
        for page in reader.pages:
            pages.append(page.extract_text() or "")
        text = "\n".join(pages)
    elif name.endswith(".docx"):
        from io import BytesIO
        document = Document(BytesIO(raw))
        paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                paragraphs.append(" | ".join(cell.text.strip() for cell in row.cells))
        text = "\n".join(paragraphs)
    elif name.endswith(".txt"):
        text = raw.decode("utf-8", errors="replace")
    else:
        raise ValueError("Unsupported file type. Upload a PDF, DOCX, or TXT file.")

    text = text.strip()
    if len(text) < 80:
        raise ValueError(
            "I couldn't extract enough text. If this is a scanned/image-only PDF, convert it to a text-based PDF or DOCX and try again."
        )
    return text[:MAX_TEXT_CHARS]


def build_prompt(resume: str, target_role: str, job_description: str, experience_level: str) -> str:
    return f'''You are a careful resume reviewer and interview coach. Review the resume below for the target role.

IMPORTANT RULES:
- This is an advisory ATS-readiness estimate, NOT a real ATS scan or a guarantee of hiring outcomes.
- Score only the information provided. Do not invent experience, qualifications, metrics, employers, dates, or skills.
- Explain that actual ATS systems vary by employer and software.
- If a job description is provided, assess keyword alignment against it; otherwise assess general resume readability and completeness for the target role.
- Recommend truthful improvements. For missing evidence, ask the candidate to add it only if accurate.
- Keep the advice practical, specific, and beginner-friendly.
- Return ONLY valid JSON matching the schema below. No markdown fences or text outside JSON.

Target role: {target_role or 'Not specified'}
Experience level: {experience_level}
Job description (may be empty):
{job_description[:12000] or 'Not provided'}

Resume text:
{resume}

Return this JSON structure:
{{
  "ats_score": 0,
  "score_label": "Short interpretation",
  "score_breakdown": [
    {{"category":"Keyword alignment","score":0,"max_score":25,"reason":"..."}},
    {{"category":"Structure and readability","score":0,"max_score":20,"reason":"..."}},
    {{"category":"Skills and qualifications","score":0,"max_score":20,"reason":"..."}},
    {{"category":"Experience and impact","score":0,"max_score":25,"reason":"..."}},
    {{"category":"Formatting and completeness","score":0,"max_score":10,"reason":"..."}}
  ],
  "top_strengths": ["..."],
  "priority_improvements": [
    {{"priority":"High","issue":"...","why_it_matters":"...","how_to_fix":"..."}}
  ],
  "missing_or_underused_keywords": ["..."],
  "section_feedback": [
    {{"section":"...","feedback":"...","suggested_action":"..."}}
  ],
  "bullet_rewrites": [
    {{"original_or_issue":"...","improved_example":"...","note":"Use only if factually accurate; replace placeholders with real details."}}
  ],
  "summary_rewrite": "A revised professional summary based only on supported resume facts; use placeholders or say what details are needed if information is insufficient.",
  "interview_prep": {{
    "likely_questions": ["..."],
    "star_story_prompts": ["..."],
    "questions_to_ask_interviewer": ["..."]
  }},
  "learning_plan": ["..."],
  "caveats": ["..." ]
}}

Scoring: ats_score must be an integer from 0 to 100. Each breakdown score must be between 0 and its max_score, and the five category scores should add to ats_score. Include 3-6 priority improvements, 3-8 keywords if relevant, 2-5 section feedback items, 2-4 bullet rewrites when possible, and 5-8 interview questions. Do not fabricate facts. If a section or detail is absent, say so plainly.'''


def call_gemini(api_key: str, prompt: str, model_name: str) -> dict[str, Any]:
    client = genai.Client(api_key=api_key)
    response = client.models.generate_content(
        model=model_name,
        contents=prompt,
        config={
            "temperature": 0.2,
            "response_mime_type": "application/json",
        },
    )
    response_text = (response.text or "").strip()
    if not response_text:
        raise RuntimeError("Gemini returned an empty response. Please try again.")
    # Defensive cleanup in case a model returns code fences despite JSON mode.
    response_text = re.sub(r"^```(?:json)?\s*|\s*```$", "", response_text, flags=re.IGNORECASE)
    try:
        data = json.loads(response_text)
    except json.JSONDecodeError as exc:
        raise RuntimeError("The AI response wasn't valid JSON. Please click Analyze again.") from exc
    if not isinstance(data, dict):
        raise RuntimeError("The AI returned an unexpected response format. Please try again.")
    return data


def safe_list(value: Any) -> list:
    return value if isinstance(value, list) else []


def render_bullets(items: list) -> None:
    if not items:
        st.write("No items returned for this section.")
        return
    for item in items:
        if isinstance(item, dict):
            title = item.get("priority") or item.get("category") or item.get("section") or item.get("issue") or ""
            if title:
                st.markdown(f"**{title}**")
            for key, label in [
                ("issue", "Issue"), ("reason", "Assessment"), ("why_it_matters", "Why it matters"),
                ("how_to_fix", "How to fix it"), ("feedback", "Feedback"),
                ("suggested_action", "Suggested action"), ("original_or_issue", "Original / issue"),
                ("improved_example", "Improved example"), ("note", "Note"),
            ]:
                val = item.get(key)
                if val:
                    st.markdown(f"- **{label}:** {val}")
            if "score" in item:
                score = item.get("score", 0)
                maximum = item.get("max_score", "?")
                st.caption(f"Score: {score}/{maximum}")
        else:
            st.markdown(f"- {item}")


st.title("📄 Resume ATS & Interview Coach")
st.write("Upload your resume for an AI-assisted ATS-readiness estimate, targeted improvements, stronger bullet examples, and interview preparation.")
st.info("The score is an estimate for guidance, not a score from a real employer ATS. Your results depend on the role and job description.")

with st.sidebar:
    st.header("Analysis settings")
    target_role = st.text_input("Target job title", placeholder="e.g., Junior Cybersecurity Analyst")
    experience_level = st.selectbox("Experience level", ["Student / Entry-level", "Early career", "Mid-level", "Senior", "Career changer"])
    model_name = st.text_input("Gemini model", value=DEFAULT_MODEL, help="Change this if your Google AI account uses a different available model name.")
    st.caption("Set GEMINI_API_KEY in Streamlit Cloud → App settings → Secrets. Never put the key in your code or GitHub.")

uploaded_file = st.file_uploader("Upload your resume", type=["pdf", "docx", "txt"], help="Maximum file size: 10 MB. Scanned PDFs may not contain extractable text.")
job_description = st.text_area("Paste the job description (recommended)", height=180, placeholder="Paste the job posting here to get a more relevant keyword match and recommendations.")

if uploaded_file:
    st.caption(f"Selected file: {uploaded_file.name} · {uploaded_file.size / 1024:.1f} KB")

analyze = st.button("Analyze resume", type="primary", disabled=uploaded_file is None, use_container_width=True)

if analyze:
    api_key = get_api_key()
    if not api_key:
        st.error("Gemini API key not found. Add GEMINI_API_KEY to your local .env/environment or Streamlit Cloud Secrets, then restart/redeploy the app.")
    elif not model_name.strip():
        st.error("Enter a Gemini model name in the sidebar.")
    else:
        try:
            with st.spinner("Reading your resume and preparing recommendations…"):
                resume_text = extract_resume_text(uploaded_file)
                result = call_gemini(
                    api_key,
                    build_prompt(resume_text, target_role.strip(), job_description.strip(), experience_level),
                    model_name.strip(),
                )
            st.session_state["resume_analysis_result"] = result
            st.session_state["resume_analysis_filename"] = uploaded_file.name
        except ValueError as exc:
            st.error(str(exc))
        except Exception as exc:
            # Do not print API keys or uploaded resume content in error messages.
            st.error(f"Analysis failed: {exc}")
            st.caption("Check that your API key is valid, billing/quota is available, and the selected model is enabled for your account.")

result = st.session_state.get("resume_analysis_result")
if result:
    st.divider()
    st.subheader(f"Results for {st.session_state.get('resume_analysis_filename', 'your resume')}")
    score = result.get("ats_score", 0)
    try:
        score = max(0, min(100, int(score)))
    except (TypeError, ValueError):
        score = 0
    c1, c2 = st.columns([1, 3])
    with c1:
        st.metric("Estimated ATS readiness", f"{score}/100")
        st.progress(score / 100)
    with c2:
        st.markdown(f"**{result.get('score_label', 'Resume review')}**")
        st.write("Use this as a checklist, not a hiring prediction. Compare your final resume against the exact job posting.")

    tabs = st.tabs(["Score breakdown", "Improvements", "Keywords & sections", "Rewrite examples", "Interview prep", "Export report"])
    with tabs[0]:
        render_bullets(safe_list(result.get("score_breakdown")))
    with tabs[1]:
        st.markdown("#### What's already working")
        render_bullets(safe_list(result.get("top_strengths")))
        st.markdown("#### Priority improvements")
        render_bullets(safe_list(result.get("priority_improvements")))
    with tabs[2]:
        st.markdown("#### Missing or underused keywords")
        render_bullets(safe_list(result.get("missing_or_underused_keywords")))
        st.markdown("#### Section-by-section feedback")
        render_bullets(safe_list(result.get("section_feedback")))
    with tabs[3]:
        st.markdown("#### Professional summary suggestion")
        st.write(result.get("summary_rewrite", "No summary suggestion returned."))
        st.markdown("#### Experience bullet examples")
        render_bullets(safe_list(result.get("bullet_rewrites")))
        st.warning("Only use suggested wording that accurately reflects your real experience. Never add invented achievements, metrics, or skills.")
    with tabs[4]:
        interview = result.get("interview_prep", {})
        if not isinstance(interview, dict):
            interview = {}
        st.markdown("#### Likely interview questions")
        render_bullets(safe_list(interview.get("likely_questions")))
        st.markdown("#### STAR story prompts")
        st.caption("Prepare examples using Situation, Task, Action, and Result. Use real examples from work, study, volunteering, or projects.")
        render_bullets(safe_list(interview.get("star_story_prompts")))
        st.markdown("#### Questions to ask the interviewer")
        render_bullets(safe_list(interview.get("questions_to_ask_interviewer")))
        st.markdown("#### Suggested learning plan")
        render_bullets(safe_list(result.get("learning_plan")))
    with tabs[5]:
        report_text = json.dumps(result, indent=2, ensure_ascii=False)
        st.download_button("Download full report (JSON)", data=report_text, file_name="resume_analysis_report.json", mime="application/json", use_container_width=True)
        readable = [
            "RESUME ATS & INTERVIEW COACH REPORT",
            f"File: {st.session_state.get('resume_analysis_filename', 'resume')}",
            f"Estimated ATS readiness: {score}/100",
            f"Interpretation: {result.get('score_label', '')}",
            "",
            "SCORE BREAKDOWN",
        ]
        for item in safe_list(result.get("score_breakdown")):
            if isinstance(item, dict):
                readable.append(f"- {item.get('category', 'Category')}: {item.get('score', '?')}/{item.get('max_score', '?')} — {item.get('reason', '')}")
        readable += ["", "PRIORITY IMPROVEMENTS"]
        for item in safe_list(result.get("priority_improvements")):
            if isinstance(item, dict):
                readable.append(f"- {item.get('priority', '')}: {item.get('issue', '')}\n  Why: {item.get('why_it_matters', '')}\n  Fix: {item.get('how_to_fix', '')}")
            else:
                readable.append(f"- {item}")
        readable += ["", "INTERVIEW PREP"]
        interview = result.get("interview_prep", {})
        if isinstance(interview, dict):
            for heading, key in [("Likely questions", "likely_questions"), ("STAR prompts", "star_story_prompts"), ("Questions for interviewer", "questions_to_ask_interviewer")]:
                readable.append(heading + ":")
                readable.extend(f"- {x}" for x in safe_list(interview.get(key)))
        readable += ["", "FULL STRUCTURED RESULTS", report_text]
        st.download_button("Download readable report (TXT)", data="\n".join(readable), file_name="resume_analysis_report.txt", mime="text/plain", use_container_width=True)

st.divider()
st.caption("Privacy note: this app sends extracted resume text and the optional job description to Google's Gemini API for analysis. Do not upload information you are not comfortable sharing with that service. This demo does not save uploaded files to a database.")
